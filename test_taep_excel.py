"""
The tariff Excel round trip and the episode export.
"""
import datetime as dt
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

import taep
from taep import CostingError


def read(data, sheet=None):
    workbook = load_workbook(BytesIO(data), data_only=False)
    return workbook[sheet] if sheet else workbook.active


def edit_download(ctx, changes=None, drop=(), add=()):
    """Download the tariff, apply edits as a human would, return the bytes."""
    data = taep.export_tariff_workbook(ctx)
    workbook = load_workbook(BytesIO(data))
    sheet = workbook[taep.TARIFF_SHEET]
    codes = {sheet.cell(r, 1).value: r for r in range(2, sheet.max_row + 1)}

    for code, column, value in (changes or []):
        # Assign through .value: openpyxl's cell(row, col, value) ignores None and
        # would silently leave the old contents in place.
        sheet.cell(codes[code], column).value = value
    for code in drop:
        sheet.delete_rows(codes[code])
        codes = {sheet.cell(r, 1).value: r for r in range(2, sheet.max_row + 1)}
    for row in add:
        sheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

def test_the_download_contains_every_tariff_row(seeded):
    sheet = read(taep.export_tariff_workbook(seeded), taep.TARIFF_SHEET)
    assert sheet.max_row == 53          # 52 tariff rows plus the header
    assert [c.value for c in sheet[1]] == [label for _, label in taep.TARIFF_COLUMNS]


def test_the_download_carries_the_greek_descriptions_and_tiers(seeded):
    sheet = read(taep.export_tariff_workbook(seeded), taep.TARIFF_SHEET)
    rows = {sheet.cell(r, 1).value: [sheet.cell(r, c).value for c in range(1, 9)]
            for r in range(2, sheet.max_row + 1)}
    assert rows["SHSO-ER11"][1] == "Καρδιογράφημα"
    assert rows["SHSO-ER11"][5] == 20.0
    assert rows["SHSO-ER16-B"][3] == "Πλύση οφθαλμού, 3 ώρες"
    assert rows["SHSO-ER39"][6] == 5.0      # the hourly rate
    assert rows["SHSO-ER3"][4] == "tariff_lookup"


def test_the_download_includes_instructions(seeded):
    workbook = load_workbook(BytesIO(taep.export_tariff_workbook(seeded)))
    assert "Οδηγίες" in workbook.sheetnames
    text = "\n".join(str(row[0].value or "") for row in workbook["Οδηγίες"].iter_rows())
    assert "όλες μαζί ή καμία" in text


# ---------------------------------------------------------------------------
# Upload validation
# ---------------------------------------------------------------------------

def test_an_unchanged_round_trip_produces_no_diff(seeded):
    rows, errors = taep.read_tariff_workbook(taep.export_tariff_workbook(seeded))
    assert errors == []
    assert len(rows) == 52
    difference = taep.diff_tariff(seeded, rows)
    assert difference == {"added": [], "changed": [], "deactivated": []}


def test_a_file_that_is_not_a_workbook_is_refused(seeded):
    rows, errors = taep.read_tariff_workbook(b"this is not excel")
    assert rows == []
    assert "έγκυρο βιβλίο Excel" in errors[0]["message"]


def test_renamed_headers_are_refused_with_a_pointer_to_the_download(seeded):
    data = edit_download(seeded)
    workbook = load_workbook(BytesIO(data))
    workbook[taep.TARIFF_SHEET].cell(1, 1, "Code")
    buffer = BytesIO()
    workbook.save(buffer)
    rows, errors = taep.read_tariff_workbook(buffer.getvalue())
    assert rows == []
    assert "Κατεβάστε ξανά" in errors[0]["message"]


@pytest.mark.parametrize("column,value,fragment", [
    (5, "nonsense", "Μη έγκυρος τύπος τιμής"),
    (6, "abc", "Μη έγκυρο βασικό ποσό"),
    (6, "-5", "αρνητικό"),
])
def test_a_bad_cell_is_reported_with_its_row_number(seeded, column, value, fragment):
    data = edit_download(seeded, changes=[("SHSO-ER11", column, value)])
    _, errors = taep.read_tariff_workbook(data)
    assert errors, "the bad cell was accepted"
    assert any(fragment in e["message"] for e in errors)
    assert all(isinstance(e["row"], int) for e in errors)


def test_a_duplicate_code_is_reported(seeded):
    data = edit_download(seeded, add=[["SHSO-ER11", "Αντίγραφο", "", "", "fixed",
                                       20, None, "ΝΑΙ"]])
    _, errors = taep.read_tariff_workbook(data)
    assert any("εμφανίζεται ξανά" in e["message"] for e in errors)


def test_fixed_plus_hourly_without_an_hourly_rate_is_refused(seeded):
    data = edit_download(seeded, changes=[("SHSO-ER39", 7, None)])
    _, errors = taep.read_tariff_workbook(data)
    assert any("απαιτεί ωριαία" in e["message"] for e in errors)


def test_a_missing_base_amount_is_refused_except_for_a_lookup(seeded):
    data = edit_download(seeded, changes=[("SHSO-ER11", 6, None)])
    _, errors = taep.read_tariff_workbook(data)
    assert any("Λείπει το βασικό ποσό" in e["message"] for e in errors)

    # tariff_lookup legitimately has no amount.
    rows, errors = taep.read_tariff_workbook(taep.export_tariff_workbook(seeded))
    lookup = next(r for r in rows if r["code"] == "SHSO-ER3")
    assert lookup["price_type"] == "tariff_lookup"
    assert errors == []


def test_blank_rows_are_ignored(seeded):
    data = edit_download(seeded, add=[[None] * 8, [None] * 8])
    rows, errors = taep.read_tariff_workbook(data)
    assert errors == []
    assert len(rows) == 52


# ---------------------------------------------------------------------------
# The diff
# ---------------------------------------------------------------------------

def test_a_price_change_shows_both_values(seeded):
    data = edit_download(seeded, changes=[("SHSO-ER11", 6, 25)])
    rows, _ = taep.read_tariff_workbook(data)
    difference = taep.diff_tariff(seeded, rows)

    assert difference["added"] == [] and difference["deactivated"] == []
    assert len(difference["changed"]) == 1
    entry = difference["changed"][0]
    assert entry["code"] == "SHSO-ER11"
    assert entry["differences"]["base_amount"] == (Decimal("20.00"), Decimal("25.00"))


def test_a_new_code_shows_as_added(seeded):
    data = edit_download(seeded, add=[["SHSO-ER99", "Νέα χρέωση", "ΑΛΛΑ", "",
                                       "fixed", 33.5, None, "ΝΑΙ"]])
    rows, errors = taep.read_tariff_workbook(data)
    assert errors == []
    difference = taep.diff_tariff(seeded, rows)
    assert [r["code"] for r in difference["added"]] == ["SHSO-ER99"]


def test_a_removed_code_shows_as_deactivated_never_deleted(seeded):
    data = edit_download(seeded, drop=["SHSO-ER11"])
    rows, _ = taep.read_tariff_workbook(data)
    difference = taep.diff_tariff(seeded, rows)
    assert [r["code"] for r in difference["deactivated"]] == ["SHSO-ER11"]


def test_a_description_change_shows_as_changed(seeded):
    data = edit_download(seeded, changes=[("SHSO-ER11", 2, "Ηλεκτροκαρδιογράφημα")])
    rows, _ = taep.read_tariff_workbook(data)
    entry = taep.diff_tariff(seeded, rows)["changed"][0]
    assert entry["differences"]["description_el"] == ("Καρδιογράφημα",
                                                      "Ηλεκτροκαρδιογράφημα")


# ---------------------------------------------------------------------------
# Applying — all of it or none of it
# ---------------------------------------------------------------------------

def test_applying_writes_adds_changes_and_deactivations(seeded):
    data = edit_download(seeded,
                         changes=[("SHSO-ER11", 6, 25)],
                         drop=["SHSO-ER14"],
                         add=[["SHSO-ER99", "Νέα χρέωση", "ΑΛΛΑ", "", "fixed",
                               33.5, None, "ΝΑΙ"]])
    rows, errors = taep.read_tariff_workbook(data)
    assert errors == []

    applied = taep.apply_tariff(seeded, rows, user_id=9)
    assert applied == {"added": 1, "changed": 1, "deactivated": 1}

    db = seeded["db_execute"]
    assert taep.money_from_db(db("SELECT base_amount FROM taep_tariff WHERE code='SHSO-ER11'",
                                fetch=True)[0]["base_amount"]) == Decimal("25.00")
    assert db("SELECT active FROM taep_tariff WHERE code='SHSO-ER14'",
              fetch=True)[0]["active"] == 0
    added = db("SELECT * FROM taep_tariff WHERE code='SHSO-ER99'", fetch=True)[0]
    assert added["description_el"] == "Νέα χρέωση"
    assert taep.money_from_db(added["base_amount"]) == Decimal("33.50")


def test_a_file_with_one_bad_row_applies_nothing(seeded):
    """Brief §7: validate before applying; never partially apply a file."""
    before = seeded["db_execute"](
        "SELECT code, base_amount, active FROM taep_tariff ORDER BY code", fetch=True)

    data = edit_download(seeded, changes=[("SHSO-ER11", 6, 25),
                                         ("SHSO-ER14", 5, "nonsense")])
    uploaded = taep.read_tariff_workbook(data)
    assert uploaded[1], "the bad row was not detected"

    with pytest.raises(CostingError) as excinfo:
        taep.apply_tariff(seeded, uploaded, user_id=9)
    assert excinfo.value.code == "TARIFF_FILE_INVALID"

    after = seeded["db_execute"](
        "SELECT code, base_amount, active FROM taep_tariff ORDER BY code", fetch=True)
    assert after == before, "a row was applied from a file that should have been rejected"


def test_an_empty_file_is_refused(seeded):
    with pytest.raises(CostingError) as excinfo:
        taep.apply_tariff(seeded, ([], []), user_id=9)
    assert excinfo.value.code == "TARIFF_FILE_EMPTY"


def test_applying_is_audit_logged(seeded):
    data = edit_download(seeded, changes=[("SHSO-ER11", 6, 25)])
    rows, _ = taep.read_tariff_workbook(data)
    taep.apply_tariff(seeded, rows, user_id=9)
    entries = [e for e in seeded["activity_log"] if e[2] == "TARIFF_UPLOADED"]
    assert len(entries) == 1 and "άλλαξαν 1" in entries[0][3]


def test_a_deactivated_tariff_can_no_longer_be_charged(seeded):
    """Deactivation is not deletion, but it must stop new charges."""
    data = edit_download(seeded, drop=["SHSO-ER11"])
    rows, _ = taep.read_tariff_workbook(data)
    taep.apply_tariff(seeded, rows, user_id=9)

    class Form:
        def getlist(self, key):
            return {"tariff_code": ["SHSO-ER11"], "tariff_quantity": ["1"]}.get(key, [])

    with pytest.raises(CostingError) as excinfo:
        taep._tariff_lines_from_form(seeded, Form())
    assert excinfo.value.code == "UNKNOWN_TARIFF"


def test_applying_twice_is_idempotent(seeded):
    data = edit_download(seeded, changes=[("SHSO-ER11", 6, 25)])
    rows, _ = taep.read_tariff_workbook(data)
    assert taep.apply_tariff(seeded, rows, user_id=9)["changed"] == 1
    assert taep.apply_tariff(seeded, rows, user_id=9) == {
        "added": 0, "changed": 0, "deactivated": 0}


# ---------------------------------------------------------------------------
# The episode export
# ---------------------------------------------------------------------------

def test_the_episode_export_has_live_formulas_not_pasted_totals(seeded):
    """An auditable workbook: the total is a formula over the rows, not a number."""
    import test_taep_episode as episodes

    for index in range(3):
        episode_id = episodes.make_episode(seeded, episode_number=f"E-{index}")
        result = episodes.price(seeded, episode_id, ["AED008", "AET057"])
        taep.finalise_episode(seeded, episode_id, result, user_id=1)

    sheet = read(taep.export_episodes_workbook(seeded, "NGH"))
    assert sheet.max_row >= 4

    formulas = [c.value for row in sheet.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=SUM(L2:L") for f in formulas), "no live total"
    assert any(f.startswith("=COUNTA(A2:A") for f in formulas), "no live count"


def test_the_episode_export_carries_the_greek_status_and_money_format(seeded):
    import test_taep_episode as episodes
    episode_id = episodes.make_episode(seeded)
    result = episodes.price(seeded, episode_id, ["AED008", "AET057"])
    taep.finalise_episode(seeded, episode_id, result, user_id=1)

    sheet = read(taep.export_episodes_workbook(seeded, "NGH"))
    values = [sheet.cell(2, column).value for column in range(1, 14)]
    assert values[0] == "OKY1054/0001"
    assert values[3] == "Παπαδόπουλος"
    assert values[9] == "ΑΙΤΗΤΕΣ ΑΣΥΛΟΥ/ΔΟΜΕΣ"
    assert values[11] == 120.0
    assert values[12] == "Οριστικοποιημένο"
    assert sheet.cell(2, 12).number_format == "#,##0.00"


def test_the_episode_export_respects_the_filters(seeded):
    import test_taep_episode as episodes
    first = episodes.make_episode(seeded, episode_number="E-1")
    taep.finalise_episode(seeded, first, episodes.price(seeded, first,
                                                        ["AED008", "AET057"]), 1)
    episodes.make_episode(seeded, episode_number="E-2", id_number="7654321")

    everything = read(taep.export_episodes_workbook(seeded, "NGH"))
    filtered = read(taep.export_episodes_workbook(seeded, "NGH",
                                                  {"status": "FINALISED"}))
    assert everything.max_row > filtered.max_row


def test_the_episode_export_is_scoped_to_the_hospital(seeded):
    import test_taep_episode as episodes
    episodes.make_episode(seeded, episode_number="E-1")
    sheet = read(taep.export_episodes_workbook(seeded, "LGH"))
    assert sheet.max_row == 1, "another hospital's episodes leaked into the export"


def test_an_empty_export_still_produces_a_usable_workbook(seeded):
    sheet = read(taep.export_episodes_workbook(seeded, "FAM"))
    assert sheet.max_row == 1
    assert sheet.cell(1, 1).value == "Αρ. Κοστολόγησης"
