"""
Phase 1 seed-data analysis for the TAEP non-GESY costing tool.

Produces, from seed/*.csv only (no hand-typed data):
  docs/overlap_decision_list.csv   machine-readable seed for service_tariff_overlap
  docs/overlap_decision_list.xlsx  decision workbook for Monada Elegchou Esodon
  stdout                           the data-quality findings

Re-run after any seed change:  python3 tools/analyse_seed.py
"""
import csv, re, sys, unicodedata, collections
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SEED, DOCS = ROOT / "seed", ROOT / "docs"
DOCS.mkdir(exist_ok=True)


def normalise(text: str) -> str:
    """Accent-, case- and final-sigma-insensitive key for Greek description matching."""
    decomposed = unicodedata.normalize("NFD", text.lower())
    stripped = "".join(c for c in decomposed if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-zα-ω0-9]+", " ", stripped.replace("ς", "σ")).strip()


def read(name):
    with open(SEED / name, encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def first_price(raw: str):
    """Leading numeric amount of a price_raw cell, or None if it carries no number."""
    match = re.search(r"\d+(?:[.,]\d+)?", raw)
    return float(match.group().replace(",", ".")) if match else None


def find_service_defects(services):
    codes = collections.Counter(s["code"] for s in services)
    duplicate_codes = {c: [s for s in services if s["code"] == c] for c, n in codes.items() if n > 1}

    by_description = collections.defaultdict(list)
    for s in services:
        by_description[normalise(s["description_el"])].append(s)
    duplicate_descriptions = {d: v for d, v in by_description.items() if len(v) > 1}

    gaps = {}
    for prefix in ("AET", "AED"):
        numbers = sorted(int(s["code"][3:]) for s in services if s["code"].startswith(prefix))
        gaps[prefix] = [f"{prefix}{n:03d}" for n in range(1, max(numbers) + 1) if n not in numbers]

    return duplicate_codes, duplicate_descriptions, gaps


def find_overlaps(services, tariffs):
    """Exact description matches between a weight service and a tariff item."""
    index = collections.defaultdict(list)
    for s in services:
        index[normalise(s["description_el"])].append(s)

    pairs, unmatched = [], []
    for t in tariffs:
        matches = index.get(normalise(t["service_el"]))
        if matches:
            pairs.extend((s, t) for s in matches)
        else:
            unmatched.append(t)
    return pairs, unmatched


def write_csv(pairs):
    path = DOCS / "overlap_decision_list.csv"
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([
            "service_code", "service_type", "service_category", "service_description_el",
            "tariff_code", "tariff_description_el", "tariff_price_raw",
            "tariff_price_parsed_eur", "proposed_treatment", "ruling", "ruled_by", "ruled_on",
        ])
        for service, tariff in pairs:
            price = first_price(tariff["price_raw"])
            writer.writerow([
                service["code"], service["service_type"], service["category"],
                service["description_el"], tariff["code"], tariff["service_el"],
                tariff["price_raw"], "" if price is None else f"{price:.2f}",
                "weight_only", "", "", "",
            ])
    return path


def write_workbook(pairs, unmatched):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    money = '#,##0.00'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Αποφάσεις επικάλυψης"

    headers = [
        "Κωδικός υπηρεσίας", "Τύπος", "Κατ.", "Περιγραφή υπηρεσίας",
        "Κωδικός τιμοκαταλόγου", "Περιγραφή τιμοκαταλόγου", "Τιμή (όπως δόθηκε)",
        "Τιμή €", "Πρόταση", "ΑΠΟΦΑΣΗ (συμπληρώστε)", "Από", "Ημερομηνία",
        "Πρόσθετη χρέωση €",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for service, tariff in sorted(pairs, key=lambda p: -(first_price(p[1]["price_raw"]) or 0)):
        price = first_price(tariff["price_raw"])
        sheet.append([
            service["code"], service["service_type"], int(service["category"]),
            service["description_el"], tariff["code"], tariff["service_el"],
            tariff["price_raw"], price, "weight_only", "", "", "",
        ])

    last = sheet.max_row
    # Exposure is live: the euro this pair adds on top of the weight cost is the tariff
    # price whenever the effective ruling is 'both' or 'tariff_only', and zero under
    # 'weight_only'. Column J overrides the proposal in column I once the Monada fills it in.
    for row in range(2, last + 1):
        sheet.cell(row=row, column=8).number_format = money
        exposure = sheet.cell(row=row, column=13)
        # Effective ruling = column J if filled in, otherwise the proposal in column I.
        exposure.value = (
            f'=IF($H{row}="","",'
            f'IF(OR(IF($J{row}="",$I{row},$J{row})="both",'
            f'IF($J{row}="",$I{row},$J{row})="tariff_only"),$H{row},0))'
        )
        exposure.number_format = money

    rule = DataValidation(
        type="list", formula1='"weight_only,tariff_only,both"',
        allow_blank=True, showDropDown=False,
        error="Επιλέξτε weight_only, tariff_only ή both.",
        errorTitle="Μη έγκυρη απόφαση",
    )
    sheet.add_data_validation(rule)
    rule.add(f"J2:J{last}")

    total = last + 2
    sheet.cell(row=total, column=9, value="Ζεύγη χωρίς απόφαση").font = Font(bold=True)
    sheet.cell(row=total, column=10, value=f'=COUNTBLANK(J2:J{last})').font = Font(bold=True)
    sheet.cell(row=total + 1, column=9,
               value="Πρόσθετη χρέωση με τις τρέχουσες αποφάσεις €").font = Font(bold=True)
    cell = sheet.cell(row=total + 1, column=10, value=f'=SUM(M2:M{last})')
    cell.font, cell.number_format = Font(bold=True), money
    sheet.cell(row=total + 2, column=9,
               value="Πρόσθετη χρέωση αν όλα κριθούν «both» €").font = Font(bold=True)
    cell = sheet.cell(row=total + 2, column=10, value=f'=SUM(H2:H{last})')
    cell.font, cell.number_format = Font(bold=True), money

    for column, width in zip("ABCDEFGHIJKLM", (13, 13, 5, 52, 12, 52, 22, 12, 13, 18, 14, 12, 13)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M{last}"

    extra = workbook.create_sheet("Πράγματι πρόσθετες")
    extra.append(["Κωδικός", "Περιγραφή", "Τιμή (όπως δόθηκε)", "Τιμή €"])
    for cell in extra[1]:
        cell.fill, cell.font = header_fill, header_font
    for tariff in unmatched:
        extra.append([tariff["code"], tariff["service_el"], tariff["price_raw"],
                      first_price(tariff["price_raw"])])
    for row in range(2, extra.max_row + 1):
        extra.cell(row=row, column=4).number_format = money
    for column, width in zip("ABCD", (13, 66, 26, 12)):
        extra.column_dimensions[column].width = width
    extra.freeze_panes = "A2"

    path = DOCS / "overlap_decision_list.xlsx"
    workbook.save(path)
    return path


def main():
    services = read("services.csv")
    tariffs = read("extra_charges.csv")
    categories = read("financial_categories.csv")

    duplicate_codes, duplicate_descriptions, gaps = find_service_defects(services)
    pairs, unmatched = find_overlaps(services, tariffs)

    print(f"services: {len(services)} rows, {len({s['code'] for s in services})} distinct codes")
    print(f"tariffs : {len(tariffs)} rows")
    print(f"\nDuplicate service codes (GO-LIVE BLOCKER): {list(duplicate_codes)}")
    for code, rows in duplicate_codes.items():
        for r in rows:
            print(f"    {code}  cat {r['category']}  {r['description_el']}")
    print("\nDuplicate descriptions under different codes:")
    for rows in duplicate_descriptions.values():
        print(f"    {[r['code'] for r in rows]}  {rows[0]['description_el'][:60]}")
    print(f"\nCode gaps: {gaps}")

    unparseable = [t for t in tariffs if not re.fullmatch(r"\s*\d+(?:[.,]\d+)?\s*", t["price_raw"])]
    print(f"\nTariff rows whose price is not a clean number: {len(unparseable)}")
    for t in unparseable:
        print(f"    {t['code']:<11} {t['price_raw'][:70]}")

    unconfirmed = [c for c in categories if c["valid_for_ae"] == "TRUE" and c["fee_status"] == "UNCONFIRMED"]
    print(f"\nTAEP-valid categories with an UNCONFIRMED registration fee: {len(unconfirmed)}")
    print("    " + ", ".join(c["code_new"] for c in unconfirmed))

    tariff_codes_overlapping = {t["code"] for _, t in pairs}
    print(f"\nOVERLAP: {len(tariff_codes_overlapping)} of {len(tariffs)} tariff items "
          f"({len(tariff_codes_overlapping) * 100 // len(tariffs)}%) are also weight services; "
          f"{len(pairs)} pairs; {len(unmatched)} genuine extras")

    print(f"\nwrote {write_csv(pairs).relative_to(ROOT)}")
    print(f"wrote {write_workbook(pairs, unmatched).relative_to(ROOT)}")

    if duplicate_codes:
        print("\nFAIL: duplicate service codes must be resolved before Phase 2.")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
