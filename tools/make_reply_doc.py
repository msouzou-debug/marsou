# -*- coding: utf-8 -*-
"""
Builds the Greek reply to the Μονάδα Ελέγχου Εσόδων for the comments of 23/10/2026.

    python3 tools/make_reply_doc.py

Renders docs/TAEP_apantiseis_23102026.md. Deliberately a small Markdown subset — the
headings, paragraphs, bullets, tables and fenced blocks this document actually uses —
rather than a general converter, so what it produces can be read off the source.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "TAEP_apantiseis_23102026.md"
TARGET = ROOT / "docs" / "TAEP_apantiseis_23102026.docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x58, 0x59, 0x5B)
GREEN = RGBColor(0x4A, 0x6B, 0x1F)


def runs(paragraph, text):
    """Bold for **…**, monospace for `…`. Nothing else, because nothing else is used."""
    for piece in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(piece)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.2)

    index = 0
    while index < len(lines):
        line = lines[index]

        if line.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.space_after = Pt(10)
            run = paragraph.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)

        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            index -= 1
            rows = [[cell.strip() for cell in row.strip("|").split("|")]
                    for row in table_lines
                    if not set(row.replace("|", "").strip()) <= set("-: ")]
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for position, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = ""
                    runs(cell.paragraphs[0], text)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        for run in paragraph.runs:
                            run.font.size = Pt(9.5)
                            if position == 0:
                                run.bold = True
                                run.font.color.rgb = NAVY

        elif line.startswith("# "):
            heading = document.add_paragraph()
            run = heading.add_run(line[2:])
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = NAVY
            heading.paragraph_format.space_after = Pt(12)

        elif line.startswith("## "):
            heading = document.add_paragraph()
            run = heading.add_run(line[3:].replace("**", ""))
            run.bold = True
            run.font.size = Pt(12.5)
            run.font.color.rgb = NAVY
            heading.paragraph_format.space_before = Pt(14)
            heading.paragraph_format.space_after = Pt(6)

        elif line.startswith("> "):
            quote = document.add_paragraph()
            quote.paragraph_format.left_indent = Cm(0.8)
            run = quote.add_run(line[2:])
            run.italic = True
            run.font.color.rgb = GREEN

        elif line.startswith("- "):
            bullet = document.add_paragraph(style="List Bullet")
            runs(bullet, line[2:])
            bullet.paragraph_format.space_after = Pt(3)

        elif line.strip() == "---":
            rule = document.add_paragraph()
            rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = rule.add_run("· · ·")
            run.font.color.rgb = GREY

        elif line.strip():
            body = document.add_paragraph()
            runs(body, line)
            body.paragraph_format.space_after = Pt(8)

        index += 1

    document.save(TARGET)
    print("wrote %s" % TARGET)
    return 0


if __name__ == "__main__":
    sys.exit(build())
