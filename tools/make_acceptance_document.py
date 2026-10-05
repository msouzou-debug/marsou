# -*- coding: utf-8 -*-
"""
Generates the acceptance document from an actual run of the scenarios.

Brief §13: the acceptance scenarios are a contractual deliverable, written as
executable tests, with the Word document generated from the test run. So this script
runs every scenario in acceptance.py against a real database and a real Flask app, and
writes what happened. A scenario that fails is reported as failed, with its error.

    python3 tools/make_acceptance_document.py

Exits non-zero if any scenario fails, so it is usable as a release gate.
"""
import sys
import traceback
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import acceptance          # noqa: E402
import taep                # noqa: E402
import taep_harness        # noqa: E402

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xA3, 0x1E, 0x1E)


def run_scenarios():
    """Run every scenario in its own clean environment. Returns a result per scenario."""
    results = []
    for scenario in acceptance.SCENARIOS:
        started = datetime.now()
        try:
            with taep_harness.Harness() as harness:
                checks = scenario.run(harness) or []
            results.append({"scenario": scenario, "passed": True,
                            "checks": list(checks), "error": None,
                            "seconds": (datetime.now() - started).total_seconds()})
        except Exception as exc:
            results.append({"scenario": scenario, "passed": False, "checks": [],
                            "error": f"{type(exc).__name__}: {exc}",
                            "trace": traceback.format_exc(limit=3),
                            "seconds": (datetime.now() - started).total_seconds()})
        print(f"  {scenario.code}  {'ΟΚ  ' if results[-1]['passed'] else 'ΑΠΟΤΥΧΙΑ'}"
              f"  {scenario.title[:62]}")
    return results


def build_document(results, run_at):
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)

    title = document.add_paragraph()
    run = title.add_run("Κοστολόγηση Περιστατικών ΤΑ.ΕΠ. — μη δικαιούχοι ΓεΣΥ")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(15), NAVY

    subtitle = document.add_paragraph()
    run = subtitle.add_run("Σενάρια Αποδοχής")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(13), NAVY

    passed = [r for r in results if r["passed"]]
    failed = [r for r in results if not r["passed"]]
    checks = sum(len(r["checks"]) for r in results)

    meta = document.add_paragraph()
    for label, value in (
        ("Ημερομηνία εκτέλεσης: ", f"{run_at:%d/%m/%Y %H:%M}\n"),
        ("Σενάρια: ", f"{len(results)} — επιτυχία {len(passed)}, "
                      f"αποτυχία {len(failed)}\n"),
        ("Έλεγχοι: ", f"{checks}\n"),
        ("Έκδοση αλγορίθμου: ", f"{taep.ALGORITHM_VERSION}\n"),
        ("Παράγεται από: ", "tools/make_acceptance_document.py"),
    ):
        meta.add_run(label).bold = True
        meta.add_run(value)

    intro = document.add_paragraph()
    intro.add_run(
        "Το παρόν δεν συντάχθηκε με το χέρι. Κάθε σενάριο είναι εκτελέσιμος κώδικας "
        "στο acceptance.py και εκτελέστηκε πραγματικά, πάνω σε βάση δεδομένων με το "
        "σχήμα και τα δεδομένα του συστήματος και πάνω στις πραγματικές οθόνες. "
        "Το αποτέλεσμα κάθε ελέγχου παρακάτω είναι αυτό που παρήγαγε η εκτέλεση. "
        "Σενάριο που αποτυγχάνει εμφανίζεται ως αποτυχία, με το σφάλμα του."
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if failed:
        warning = document.add_paragraph()
        count = len(failed)
        run = warning.add_run(
            f"ΠΡΟΣΟΧΗ: {count} "
            f"{'σενάριο απέτυχε' if count == 1 else 'σενάρια απέτυχαν'}. "
            f"Το σύστημα δεν είναι έτοιμο για παραγωγική λειτουργία.")
        run.bold, run.font.color.rgb = True, RED

    # Summary table first: the page a reader signs off from.
    document.add_paragraph()
    heading = document.add_paragraph()
    run = heading.add_run("Σύνοψη")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(12), NAVY

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells,
                           ("Κωδ.", "Σενάριο", "Παράγραφος", "Αποτέλεσμα")):
        cell.text = ""
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
    for result in results:
        row = table.add_row().cells
        row[0].text = result["scenario"].code
        row[1].text = result["scenario"].title
        row[2].text = result["scenario"].clause
        run = row[3].paragraphs[0].add_run("Επιτυχία" if result["passed"]
                                          else "ΑΠΟΤΥΧΙΑ")
        run.bold = True
        run.font.color.rgb = GREEN if result["passed"] else RED
    for width, column in zip((1.6, 9.0, 2.2, 2.4), range(4)):
        for row in table.rows:
            row.cells[column].width = Cm(width)

    document.add_page_break()
    heading = document.add_paragraph()
    run = heading.add_run("Αναλυτικά σενάρια")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(12), NAVY

    for result in results:
        scenario = result["scenario"]
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"{scenario.code} — {scenario.title}")
        run.bold, run.font.size = True, Pt(11)
        run.font.color.rgb = GREEN if result["passed"] else RED

        for label, text in (("Δεδομένου: ", scenario.given),
                            ("Όταν: ", scenario.when),
                            ("Τότε: ", scenario.then)):
            line = document.add_paragraph()
            line.add_run(label).bold = True
            line.add_run(text)
            line.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            line.paragraph_format.space_after = Pt(2)

        reference = document.add_paragraph()
        reference.add_run("Παράγραφος προδιαγραφής: ").bold = True
        reference.add_run(scenario.clause)
        reference.paragraph_format.space_after = Pt(3)

        if result["passed"]:
            verified = document.add_paragraph()
            run = verified.add_run(f"Επαληθεύτηκε ({len(result['checks'])} έλεγχοι):")
            run.bold, run.font.color.rgb = True, GREEN
            for check in result["checks"]:
                bullet = document.add_paragraph(check, style="List Bullet")
                bullet.paragraph_format.space_after = Pt(1)
        else:
            failure = document.add_paragraph()
            run = failure.add_run("ΑΠΟΤΥΧΙΑ: ")
            run.bold, run.font.color.rgb = True, RED
            failure.add_run(result["error"])
            trace = document.add_paragraph(result.get("trace", ""))
            trace.style = document.styles["Normal"]
            trace.runs[0].font.size = Pt(7) if trace.runs else None

        document.add_paragraph().paragraph_format.space_after = Pt(8)

    document.add_page_break()
    heading = document.add_paragraph()
    run = heading.add_run("Παραλαβή")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(12), NAVY
    document.add_paragraph(
        "Με την υπογραφή του παρόντος επιβεβαιώνεται ότι τα σενάρια αποδοχής "
        "εκτελέστηκαν και ελέγχθηκαν.")
    for role in ("Μονάδα Ελέγχου Εσόδων", "Τμήμα Πληροφορικής",
                 "Διεύθυνση Νοσηλευτηρίου"):
        line = document.add_paragraph()
        line.add_run(f"{role}: ").bold = True
        line.add_run("……………………………………………   Ημερομηνία: ………………………")
        line.paragraph_format.space_after = Pt(14)

    return document


def main():
    run_at = datetime.now()
    print(f"Εκτέλεση {len(acceptance.SCENARIOS)} σεναρίων αποδοχής...")
    results = run_scenarios()

    document = build_document(results, run_at)
    path = ROOT / "docs" / "TAEP_senaria_apodochis.docx"
    document.save(path)

    failed = [r for r in results if not r["passed"]]
    checks = sum(len(r["checks"]) for r in results)
    print(f"\n{len(results) - len(failed)}/{len(results)} σενάρια, "
          f"{checks} έλεγχοι")
    print(f"wrote {path.relative_to(ROOT)}")
    if failed:
        print(f"\nΑΠΟΤΥΧΙΑ σε {len(failed)}: "
              f"{', '.join(r['scenario'].code for r in failed)}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
