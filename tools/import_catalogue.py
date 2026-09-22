# -*- coding: utf-8 -*-
"""
Rebuilds seed/services.csv and seed/care_levels.csv from the authoritative
A&E catalogue workbook.

The CSV we worked from in Phase 1 was a corrupted extract: it had lost five
treatment codes, mangled AET091 into a duplicate of AET085, and split AET092
across two rows. Four of the Phase 1 "defects" were artifacts of that file.
The workbook in seed/source/ is the source of truth; nothing is retyped.

    python3 tools/import_catalogue.py
"""
import csv
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SEED = ROOT / "seed"
WORKBOOK = SEED / "source" / "20230706_AE_Catalogue_5.xlsx"

SHEETS = {"AE Investigation": "INVESTIGATION", "AE Treatment": "TREATMENT"}
CATEGORY_LIMIT = {"INVESTIGATION": 3, "TREATMENT": 5}


def sheet_rows(workbook, name):
    """Yield rows keyed by header. The Care Levels sheet calls its code column
    "CPT Code" while the service sheets call it "Code"; normalise to "Code"."""
    sheet = workbook[name]
    headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
    headers = ["Code" if h == "CPT Code" else h for h in headers]
    for r in range(2, sheet.max_row + 1):
        row = {headers[c - 1]: sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)}
        if row.get("Code"):
            yield row


def category_number(category_code, service_type):
    """AECAT1..AECAT5 -> 1..5, checked against what the side may carry."""
    if not str(category_code or "").startswith("AECAT"):
        raise SystemExit(f"Unexpected category code {category_code!r}")
    value = int(str(category_code)[5:])
    limit = CATEGORY_LIMIT[service_type]
    if not 1 <= value <= limit:
        raise SystemExit(f"{service_type}: category {value} exceeds the permitted {limit}")
    return value


def main():
    workbook = load_workbook(WORKBOOK, data_only=True)

    services = []
    for sheet_name, service_type in SHEETS.items():
        for row in sheet_rows(workbook, sheet_name):
            services.append({
                "code": str(row["Code"]).strip(),
                "service_type": service_type,
                "category": category_number(row["Category"], service_type),
                "category_code": str(row["Category"]).strip(),
                "description_el": str(row["Medium Description (GR)"] or "").strip(),
                "description_en": str(row["Medium Description (EN)"] or "").strip(),
            })

    codes = [s["code"] for s in services]
    if len(codes) != len(set(codes)):
        raise SystemExit("Duplicate codes in the source workbook — stopping.")

    services.sort(key=lambda s: s["code"])
    path = SEED / "services.csv"
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(services[0]))
        writer.writeheader()
        writer.writerows(services)

    # Care levels carry the Greek band labels we refused to invent in Phase 1.
    levels = []
    for row in sheet_rows(workbook, "Care Levels"):
        levels.append({
            "code": str(row["Code"]).strip(),
            "level_code": str(row["Category"]).strip(),
            # The source has a stray double space in the AELEVEL3 Greek text.
            "label_en": " ".join(str(row["Medium Description (EN)"] or "").split()),
            "label_el": " ".join(str(row["Medium Description (GR)"] or "").split()),
        })
    levels_path = SEED / "care_levels.csv"
    with open(levels_path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(levels[0]))
        writer.writeheader()
        writer.writerows(levels)

    investigations = [s for s in services if s["service_type"] == "INVESTIGATION"]
    treatments = [s for s in services if s["service_type"] == "TREATMENT"]
    print(f"wrote {path.relative_to(ROOT)}: {len(services)} services "
          f"({len(investigations)} investigations, {len(treatments)} treatments)")
    print(f"wrote {levels_path.relative_to(ROOT)}: {len(levels)} care levels")
    for level in levels:
        print(f"    {level['level_code']:<10} {level['label_el']}")


if __name__ == "__main__":
    main()
